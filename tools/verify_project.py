"""Consistency checks for the L1/L2 deliverables, not Cisco simulation."""
from pathlib import Path
from collections import Counter
import json,math,zipfile,re
from docx import Document
from pptx import Presentation
R=Path(__file__).resolve().parents[1];a=json.loads((R/'docs/project-data.json').read_text());e=a['endpoints'];p=a['ports']
assert len(e)==len(p)==99
assert Counter(x['type'] for x in e)=={'pc':84,'printer':8,'ap':4,'server':3}
assert len({x['id'] for x in e})==99
assert {x['id'] for x in e}=={x['endpoint'] for x in p}
assert len({(x['device'],x['port']) for x in p})==99
assert all(int(x['port'].split('/')[-1])<=24 for x in p)
assert sum(x['length_m']>0 for x in e)==96
assert all(x['length_m']==0 for x in e if x['type']=='server')
assert max(x['length_m'] for x in e)+4<=100
assert math.isclose(a['copper_m'],sum(x['length_m'] for x in e))
assert a['spools']*305>=a['copper_m']*1.2
assert a['total_rub']==sum(x[2]*x[3] for x in a['equipment'])
assert [round(sum(a['loads'][f])*1.5,1) for f in ['C1','C2','B1','B2']]==[180.6,132.6,92.7,92.7]
D=Document(R/'docs/Расчёты_вариант_8.docx');S=Presentation(R/'docs/Презентация_вариант_8.pptx')
assert len(S.slides)==18
text='\n'.join(shape.text for s in S.slides for shape in s.shapes if shape.has_text_frame)
doc='\n'.join(x.text for x in D.paragraphs)+'\n'+'\n'.join(c.text for t in D.tables for row in t.rows for c in row.cells)
assert 'Устинов Илья Максимович' in text and 'ИУ10-59' in text and 'Федорова В. А.' in text
for i in range(1,15):assert any(s.shapes[0].text.startswith(str(i)+'.') for s in S.slides)
assert f'Общая длина Cat6: {a["copper_m"]:.1f} м' in doc
assert not re.search(r'\b(?:\d{1,3}\.){3}\d{1,3}\b',text+doc)
assert not any(v in (text+doc).lower() for v in ['доллар','курс валют','http://','https://'])
assert not re.search(r'\bпользователь(?:ем|я|ю)?\b', (text+doc).lower())
for f in [R/'docs/Расчёты_вариант_8.docx',R/'docs/Презентация_вариант_8.pptx']:
 with zipfile.ZipFile(f) as z:assert z.testzip() is None
print('PASS: 99 подключений, 96 постоянных линий, нагрузка, смета, 18 слайдов / 14 пунктов; только рубли, без IP-адресов.')
