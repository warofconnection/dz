"""Render the separate defense study materials. Submission files are untouched."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1]
BLUE='176DEA';GRAY='586475'
def render(source):
 d=Document();sec=d.sections[0];sec.page_height=Cm(29.7);sec.page_width=Cm(21)
 sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=sec.right_margin=Cm(2)
 styles=d.styles
 for name in ['Normal','Body Text']:
  styles[name].font.name='DejaVu Sans';styles[name].font.size=Pt(10.5)
  styles[name].paragraph_format.space_after=Pt(6)
  styles[name].paragraph_format.line_spacing=1.12
 for name,size in [('Title',25),('Heading 1',17),('Heading 2',13)]:
  styles[name].font.name='DejaVu Sans';styles[name].font.size=Pt(size);styles[name].font.color.rgb=RGBColor.from_string(BLUE)
  styles[name].paragraph_format.keep_with_next=True
 # Footer marks the file as study material, not an additional submission report.
 footer=sec.footer.paragraphs[0];footer.text='Подготовка к защите • ИТС • Вариант 8     |     '
 footer.style=styles['Caption'];footer.runs[0].font.size=Pt(8)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
 lines=source.read_text().splitlines();i=0
 while i<len(lines):
  s=lines[i].strip()
  if not s:i+=1;continue
  if s.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    row=[c.strip() for c in lines[i].strip().strip('|').split('|')]
    if not all(re.fullmatch(r'[-: ]+',c) for c in row):rows.append(row)
    i+=1
   table=d.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
   for k,row in enumerate(rows):
    cells=table.add_row().cells
    for cell,text in zip(cells,row):
     cell.text=text
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.05
      for run in p.runs:run.font.name='DejaVu Sans';run.font.size=Pt(9);run.bold=k==0
    trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    if k==0:
     repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
   d.add_paragraph();continue
  if s.startswith('!['):
   match=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',s);assert match
   path=(source.parent/match[2]).resolve();d.add_picture(str(path),width=Cm(17))
   cap=d.add_paragraph(match[1],style='Caption');cap.paragraph_format.space_after=Pt(10)
  elif s.startswith('# '):d.add_heading(s[2:],0)
  elif s.startswith('## '):d.add_heading(s[3:],1)
  elif s.startswith('### '):d.add_heading(s[4:],2)
  else:d.add_paragraph(s)
  i+=1
 d.core_properties.title=lines[0].lstrip('# ');d.core_properties.subject='Подготовка к защите проекта L1–L2 городской поликлиники';d.core_properties.author=''
 out=source.with_suffix('.docx');d.save(out);print(out.name)
for name in ['Пособие_для_защиты_ИТС_вариант8.md','Памятка_перед_защитой.md']:render(R/'study'/name)
