from pathlib import Path
import json, math, csv, html
from copy import deepcopy
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
ROOT=Path(__file__).resolve().parents[1]
FONT='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
def picture(name,w=1600,h=1000):
 im=Image.new('RGB',(w,h),'white'); return im,ImageDraw.Draw(im)
def txt(d,xy,s,size=22,fill='#172c47'):
 d.text(xy,s,font=ImageFont.truetype(FONT,size),fill=fill)
def save(im,name): im.save(ROOT/'diagrams'/f'{name}.png')
floors=[('C1','Центральный корпус — этаж 1',24,['Регистратура','Регистратура','Терапия','Терапия','Диагностика','Диагностика']),('C2','Центральный корпус — этаж 2',24,['Специалисты','Специалисты','Специалисты','Специалисты','Администрация','Администрация / ИТ']),('B1','Филиал 1 — этаж 1',18,['Регистратура','Регистратура','Терапия','Терапия','Терапия','Сестринская']),('B2','Филиал 2 — этаж 1',18,['Регистратура','Регистратура','Терапия','Терапия','Терапия','Сестринская'])]
ends=[]
for code,title,n,deps in floors:
 im,d=picture(code,1660,970); txt(d,(45,25),title,32); txt(d,(45,75),'Проектный план: 36 × 16 м; кабинеты 6 × 6 м; коридор 36 × 4 м',23)
 ox,oy,sc=80,165,40
 def pos(x,y):return (ox+x*sc,oy+y*sc)
 def room(x,y,w,h,label):
  d.rectangle([pos(x,y),pos(x+w,y+h)],outline='#172c47',width=3); txt(d,(pos(x,y)[0]+8,pos(x,y)[1]+8),label,17)
 for j in range(6):room(j*6,0,6,6,f'{j+1:02d}  {deps[j]}\n6 × 6 м')
 room(0,6,36,4,'Коридор 36 × 4 м; кабельный лоток на высоте 2,8 м')
 room(0,10,6,6,'Техническая\n6 × 6 м')
 room(6,10,30,6,'Ожидание, санитарные и вспомогательные помещения — 30 × 6 м')
 rack=(3,12); d.rectangle([pos(2.3,11.3),pos(3.7,12.7)],fill='#172c47'); txt(d,(pos(0.3,14)[0],pos(0.3,14)[1]),'Стойка R + SW',16)
 counts=[4]*6 if n==24 else [2,2,3,3,4,4]
 index=0
 def endpoint(label,typ,x,y,vlan,dep):
  global ends
  length=abs(x-3)+abs(y-8)+4+5 if typ!='server' else abs(x-3)+abs(y-12)+5
  ends.append(dict(floor=code,id=label,type=typ,x=x,y=y,vlan=vlan,department=dep,length_m=round(length,2)))
  if typ!='server':d.line([pos(*rack),pos(3,8),pos(x,8),pos(x,y)],fill='#bdd7e9',width=2)
  px,py=pos(x,y); d.rectangle((px-8,py-8,px+8,py+8),fill={'pc':'#177b9d','printer':'#9a5c22','ap':'#3f8c50','server':'#77469a'}[typ]);txt(d,(px-22,py+10),label,12)
 for j,count in enumerate(counts):
  for k in range(count):
   index+=1; x=j*6+[1.2,4.8,1.2,4.8][k];y=[2,2,4,4][k]
   vlan=20 if code=='C1' and j<2 or code=='C2' and j>=4 else 10
   endpoint(f'{code}-PC{index:02d}','pc',x,y,vlan,deps[j])
 for k,x in enumerate([8,28]):endpoint(f'{code}-PR{k+1}','printer',x,8,10,'Общий печатный ресурс')
 endpoint(f'{code}-AP1','ap',18,8,90,'Гостевой доступ')
 if code=='C1':
  for k,x in enumerate([1,3,5]):endpoint(['SRV-MIS','SRV-MAIL','SRV-FILE'][k],'server',x,13,30,'Серверная')
 txt(d,(80,840),'ПК — синий; принтер — коричневый; точка доступа — зелёный; сервер — фиолетовый.',19)
 txt(d,(80,880),'Трассы показаны по коридорному лотку; вводы в кабинеты — по стенам. Высота этажа — 3,6 м.',19)
 save(im,f'floor-{code}')
# Geographic assumptions
im,d=picture('site',1600,760)
for x,y,label in [(600,100,'Центральный корпус\n2 этажа, 36 × 16 м'),(120,480,'Филиал 1\n1 этаж, 36 × 16 м'),(1080,480,'Филиал 2\n1 этаж, 36 × 16 м')]:
 d.rounded_rectangle((x,y,x+390,y+160),radius=18,fill='#e4eef5',outline='#172c47',width=3);txt(d,(x+20,y+30),label,25)
d.line((720,260,330,480),fill='#177b9d',width=5);d.line((880,260,1250,480),fill='#177b9d',width=5)
txt(d,(270,300),'5 км; канал оператора 100 Мбит/с',22);txt(d,(920,300),'8 км; 100 Мбит/с',22)
txt(d,(180,685),'Расстояние между филиалами — 9 км. Схематично, без привязки к реальной карте.',22);save(im,'site')
im,d=picture('org',1700,950)
boxes=[(560,25,570,95,'Главный врач — 1 (в составе администрации)'),(40,220,760,280,'Центральный корпус — 64 сотрудника\nРегистратура: 10; терапия: 12; диагностика: 12\nСпециалисты: 22; администрация: 6; ИТ: 2\n48 рабочих станций'),(880,220,760,280,'Филиалы — по 24 сотрудника\nРегистратура: 6; терапия: 12; сестринская: 6\nПо 18 рабочих станций; всего 48 сотрудников'),(260,690,1180,150,'Совместная работа: регистратура → врачи → диагностика\nАдминистрация получает отчёты; ИТ обслуживает все подразделения.\nИтого: 112 сотрудников; 84 рабочие станции (сменная работа).')]
for x,y,w,h,s in boxes:d.rounded_rectangle((x,y,x+w,y+h),18,fill='#e4eef5',outline='#172c47',width=3);txt(d,(x+20,y+20),s,25)
d.line((845,120,410,220),fill='#172c47',width=3);d.line((845,120,1260,220),fill='#172c47',width=3);d.line((410,500,600,690),fill='#177b9d',width=3);d.line((1260,500,1100,690),fill='#177b9d',width=3);save(im,'org')
# Logical design illustration, explicitly not a Cisco screenshot
im,d=picture('logical',1800,1150)
txt(d,(40,20),'Проект логической схемы — эскиз для сборки в Cisco Packet Tracer',28)
boxes={ 'PUBLIC':(1440,110,'PUBLIC-WEB\n203.0.113.10'), 'ISP':(1170,110,'R-ISP\n2911'), 'WAN':(780,110,'SW-PROVIDER\nVLAN 901/902/903'), 'EDGE':(780,390,'R-EDGE\n2911, NAT'), 'C1':(310,390,'R-C1\n2911'), 'C2':(310,750,'R-C2\n2911'), 'B1':(80,110,'R-B1\n2911'), 'B2':(420,110,'R-B2\n2911') }
for a,b,label in [('PUBLIC','ISP','Gi0/1'),('ISP','WAN','Gi0/0 · VLAN 903'),('WAN','EDGE','Gi0/1 · trunk'),('C1','EDGE','Gi0/2 ↔ Gi0/0'),('C1','C2','Gi0/1 ↔ Gi0/1'),('B1','WAN','Gi0/1 · VLAN 901'),('B2','WAN','Gi0/1 · VLAN 902')]:
 x,y,_=boxes[a]; xx,yy,_=boxes[b];d.line((x+130,y+55,xx+130,yy+55),fill='#69859b',width=4);txt(d,((x+xx)//2+80,(y+yy)//2+35),label,15)
for name,(x,y,s) in boxes.items():d.rounded_rectangle((x,y,x+260,y+110),12,fill='#e4eef5',outline='#172c47',width=3);txt(d,(x+12,y+18),s,21)
for x,y,s in [(30,270,'B1: SW-B1 · 18 ПК\n2 принтера, AP'),(1120,390,'B2: SW-B2 · 18 ПК\n2 принтера, AP'),(20,560,'C1: SW-C1A + SW-C1B\n24 ПК, 2 принтера, AP\n3 сервера'),(20,920,'C2: SW-C2A + SW-C2B\n24 ПК, 2 принтера, AP')]:
 d.rectangle((x,y,x+600,y+120),outline='#177b9d',width=3);txt(d,(x+15,y+15),s,22)
txt(d,(700,820),'Межэтажные и межкорпусные связи —\nчерез маршрутизаторы. На каждом этаже —\nкоммутаторы доступа. Гости изолированы.\n\nSW-PROVIDER и R-ISP — оборудование\nоператора, в смету поликлиники не входят.',23)
save(im,'logical-draft')
import subprocess
subprocess.run(['dot','-Tpng','-Gdpi=130',str(ROOT/'packet-tracer/topology.dot'),'-o',str(ROOT/'diagrams/logical-draft.png')],check=True)
# Port assignments
ports=[]
for floor,_,_,_ in floors:
 items=[e for e in ends if e['floor']==floor]
 switches=[f'SW-{floor}A',f'SW-{floor}B'] if floor.startswith('C') else [f'SW-{floor}']
 for j,e in enumerate(items):
  sw=switches[0] if len(switches)==1 or j<12 else switches[1]
  p=j+1 if len(switches)==1 or j<12 else j-11
  ports.append(dict(device=sw,port=f'FastEthernet0/{p}',endpoint=e['id'],vlan=e['vlan']))
with (ROOT/'packet-tracer/ports.csv').open('w') as f:
 w=csv.DictWriter(f,fieldnames=ports[0].keys());w.writeheader();w.writerows(ports)
with (ROOT/'docs/cable-routes.csv').open('w') as f:
 w=csv.DictWriter(f,fieldnames=ends[0].keys());w.writeheader();w.writerows(ends)
loads={ 'C1':[9.6,12,12.8,12,64,10], 'C2':[9.6,12,12.8,12,32,10], 'B1':[7.2,10,9.6,9,16,10], 'B2':[7.2,10,9.6,9,16,10] }
equipment=[('Рабочая станция','Core i5 / Ryzen 5, 16 ГБ, SSD 512 ГБ, монитор, 1GbE',84,65000),('Сервер','8 ядер, ECC 64 ГБ, зеркальные SSD, 2 × 1GbE, ИБП отдельно',3,260000),('Маршрутизатор','Cisco C1111-8P; PT: 2911; производительность проверить с лицензией',5,180000),('Коммутатор','Cisco CBS350-24P-4G, 24 × 1GbE PoE+, 4 × SFP 1G; PT: 2960-24TT',6,95000),('Точка доступа','TP-Link EAP610, Wi-Fi 6, 1GbE, PoE; PT: AccessPoint-PT',4,16000),('Сетевое МФУ','Ethernet 1G/100M, лазерное',8,30000),('Шкаф 19 дюймов','Один на этаж; центральная серверная — полноразмерный',4,35000),('ИБП','3 серверных + 4 для этажных стоек',7,45000),('Патч-панель Cat6','24 порта',6,6000),('Розетка RJ-45 Cat6','Один порт на 96 линий, серверы в стойке',96,600),('Патч-корд Cat6','Два на каждую из 99 линий + 7 межсетевых',205,600)]
copper=sum(e['length_m'] for e in ends);spools=math.ceil(copper*1.2/305)
equipment.extend([('Кабель Cat6 U/UTP','305 м в бухте, сплошная медь, LSZH',spools,22000),('Оптический кабель OM3','Межэтажный участок: 12 м + 20% резерв → 15 м',15,180),('SFP 1000BASE-SX','Два для межэтажной линии (реальные маршрутизаторы)',2,7000),('Оптический кросс','2 оконечных комплекта, включая пигтейлы и адаптеры',2,7000),('Оптический патч-корд','LC–LC OM3, 2 м',2,1800),('Кабельный лоток','36 м × 4 этажа + 20% = 173 м',173,450)])
total=sum(n*p for _,_,n,p in equipment)
budget_text=f'{total:,}'.replace(',',' ')
project=dict(endpoints=ends,ports=ports,loads=loads,copper_m=copper,spools=spools,equipment=equipment,total_rub=total)
(ROOT/'docs/project-data.json').write_text(json.dumps(project,ensure_ascii=False,indent=2))
# DOCX
D=Document();sec=D.sections[0];sec.top_margin=sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2)
style=D.styles['Normal'];style.font.name='Times New Roman';style.font.size=Pt(12);style.paragraph_format.space_after=Pt(6)
def p(s):D.add_paragraph(s)
def h(s,l=1):D.add_heading(s,l)
def table(head,rows):
 t=D.add_table(rows=1, cols=len(head));t.style='Table Grid'
 for c,s in zip(t.rows[0].cells,head):c.text=str(s)
 for row in rows:
  for c,s in zip(t.add_row().cells,row):c.text=str(s)
h('Домашнее задание № 1. ИТС. Вариант 8')
p('Предварительное проектирование физического и канального уровней сети городской поликлиники. Расчётное приложение к презентации.')
p('Расчётное приложение содержит исходные данные, оценку сетевой нагрузки, кабельные трассы, смету и результаты моделирования в Cisco Packet Tracer 8.2.2. Цены приняты ориентировочно для оценки стоимости проекта.')
h('1. Исходные данные и границы проекта')
p('Учебная поликлиника «Городская поликлиника № 8». Центральный корпус: 2 этажа; два одноэтажных филиала. Каждый этаж: 36 × 16 м, высота 3,6 м. Филиалы удалены от центрального корпуса на 5 и 8 км, друг от друга на 9 км. Эти значения и численность персонала приняты проектировщиком, а не заданы методичкой.')
table(['Подразделение','Сотрудников','ПК','Размещение'],[['Регистратура центра',10,8,'C1, кабинеты 01–02'],['Терапия центра',12,8,'C1, 03–04'],['Диагностика',12,8,'C1, 05–06'],['Врачи-специалисты',22,16,'C2, 01–04'],['Администрация',6,6,'C2, 05 и часть 06'],['ИТ',2,2,'C2, 06'],['Каждый филиал: регистратура',6,4,'01–02'],['Каждый филиал: терапия',12,10,'03–05'],['Каждый филиал: сестринская',6,4,'06'],['Итого',112,84,'4 этажа в 3 зданиях']])
p('Рабочие места совместно используются сотрудниками разных смен; 112 сотрудников не означают 112 одновременно работающих ПК. Сеть содержит 99 оконечных проводных подключений: 84 ПК, 8 МФУ, 4 точки доступа и 3 сервера. Резерв расширения не менее 20%.')
h('2. Прикладные задачи и программное обеспечение')
p('Медицинская информационная система (МИС): запись пациентов, электронная карта, направления и результаты. Диагностика: обмен исследованиями DICOM и архив изображений PACS. Общие службы: корпоративная почта, документы, печать, Интернет и гостевой беспроводной доступ. Реальные персональные данные в учебной модели не используются.')
p('ПК: поддерживаемая ОС Windows 11 Pro либо Linux при совместимости с МИС, браузер, офисный пакет, почтовый клиент, антивирус/EDR; диагностические станции — DICOM-просмотрщик. SRV-MIS: Linux LTS, PostgreSQL, веб-приложение МИС; SRV-MAIL: SMTP/IMAP, DNS; SRV-FILE: файловая служба, PACS и резервное копирование. В Packet Tracer используются только упрощённые DNS, HTTP и EMAIL. PostgreSQL, PACS, EDR и полноценная защита медицинских данных этим симулятором не проверяются.')
h('3. Расчёт максимальной нагрузки (пункт 7 отчёта)')
p('Часы наибольшей нагрузки: будние дни 08:00–11:00 и 16:00–18:00; все указанные активные потоки считаются одновременными. Десятичные единицы: 1 Мбайт = 10⁶ байт, 1 Мбит/с = 10⁶ бит/с. Для передачи блока R = 8 × V × n / t, где V — Мбайт, n — число одновременных передач, t — допустимое время в секундах. Для фиксированных потоков R = n × r. Расчётная скорость C = 1,25 × 1,20 × ΣR = 1,50 × ΣR. 25% — условная надбавка на протоколы и вариацию; 20% — рост нагрузки. Это оценка выбранного пикового сценария, а не измеренный абсолютный максимум.')
table(['Сервис','C1','C2','Каждый филиал','Принятый профиль'],[['МИС',9.6,9.6,7.2,'12/12/9 активных × 0,8 Мбит/с'],['Почта',12,12,10,'6/6/5 сообщений × 0,5 Мбайт / 2 с × 8'],['Файлы',12.8,12.8,9.6,'4/4/3 передачи × 2 Мбайт / 5 с × 8'],['Интернет',12,12,9,'12/12/9 активных × 1 Мбит/с'],['DICOM',64,32,16,'4/2/1 передачи × 20 Мбайт / 10 с × 8'],['Гости',10,10,10,'10 клиентов на этаж × 1 Мбит/с']])
table(['Участок','Исходная сумма, Мбит/с','С резервами','Принято'],[[k,round(sum(v),2),round(sum(v)*1.5,2),'1 Гбит/с на этаж; филиал: 100 Мбит/с'] for k,v in loads.items()])
p('C2 → C1: 132,6 Мбит/с; каждый филиал → центр: 92,7 Мбит/с (консервативно весь трафик этажа проходит через канал). Нагрузка общего стыка C1 → R-EDGE: 1,5 × (88,4 + 2 × 61,8 + 22) = 351,0 Мбит/с, где 22 = Интернет и гости C1. Интернет: 1,5 × (12 + 12 + 9 + 9 + 4 × 10) = 123 Мбит/с → тариф 200 Мбит/с. Каналы филиалов по 100 Мбит/с удовлетворяют расчёту 92,7 Мбит/с с уже включёнными резервами, но имеют лишь 7,3 Мбит/с остатка; при росте DICOM следующий тариф — 200 Мбит/с.')
p('Консервативная оценка серверного стыка: 1,5 × [МИС 33,6 + почта 44 + файлы 44,8 + DICOM 128] = 375,6 Мбит/с. Для SRV-MIS: 50,4; SRV-MAIL: 66; SRV-FILE/PACS: 259,2 Мбит/с. Каждый сервер получает 1 Гбит/с. Суммарная нагрузка этажных направлений 1,5 × 332,4 = 498,6 Мбит/с используется только как верхняя оценка общего объёма; потоки не суммируются повторно на конкретном стыке.')
p('ПК диагностики: МИС 0,8 + почта 2 + файлы 3,2 + Интернет 1 + DICOM 16 = 23 Мбит/с; с резервами 34,5 Мбит/с. Остальные ПК — до 10,5 Мбит/с. Поэтому Fast Ethernet портов учебного 2960 достаточен для этого сценария, а в реальном проекте выбран 1000BASE-T для развития. Магистральные порты симулятора — GigabitEthernet.')
h('4. Стандарт, топология и оборудование (пункты 8–9)')
p('Иерархическая звезда: отдельная коммутация на каждом этаже, маршрутизация между этажами и площадками. Рабочие линии: 1000BASE-T, Cat6 U/UTP LSZH, цельномедный проводник. Постоянная линия ≤ 90 м, канал с патч-кордами ≤ 100 м. Межэтажная линия: 1000BASE-SX, OM3, 850 нм, два волокна; расчётный участок 12 м. Межрайонные каналы — арендуемые операторские Ethernet E-Line, а не собственная прокладка кабеля по городу.')
p('Коммутаторы: на каждом этаже центра по 2 × 24 порта, в каждом филиале 1 × 24. Всего 144 пользовательских порта, занято 99; резерв 45. C1: 30/48, C2: 27/48, B1 и B2: 21/24. В филиале 18 ПК + ceil(18 × 0,2) = 22 ПК; вместе с 3 прочими устройствами потребуется 25 портов: при полном 20% росте понадобится дополнительный коммутатор. На начальном этапе физический резерв филиала — 3/21 = 14,3%, что превышает минимум 5–10% методички. Резерв по нагрузке и резерв портов различаются.')
p('Маршрутизаторы поликлиники: R-C1, R-C2, R-B1, R-B2 (по этажам) и R-EDGE (стык оператора, Интернет, NAT). R-ISP и SW-PROVIDER служат только для моделирования сети оператора и не включаются в закупку. Реальные модели и учебные эквиваленты отличаются; 2911 и 2960 не предлагаются к закупке. Перед закупкой C1111-8P необходимо подтвердить скорость маршрутизации, наличие нужных SFP/медных интерфейсов и лицензий в выбранной комплектации; производительность с VPN не считается проверенной.')
h('5. Подробный расчёт кабеля (пункт 11)')
p('Начало координат каждого плана — левый верхний угол, размеры в метрах. Стойка (3;12), кабельный лоток по оси y = 8. Для каждого рабочего устройства длина L = |x−3| + |y−8| + 4 + 5; 4 м — ввод от стойки к коридору; 5 м — суммарные вертикальные спуски и сервисные петли. Для серверов внутри стойки L = |x−3| + |y−12| + 5. Все исходные координаты и индивидуальные длины приведены ниже и в cable-routes.csv. Длина не является прямой диагональю между устройствами.')
table(['Этаж','Линий','Сумма трасс, м','Максимальная линия, м'],[[code,len([e for e in ends if e['floor']==code]),round(sum(e['length_m'] for e in ends if e['floor']==code),2),max(e['length_m'] for e in ends if e['floor']==code)] for code,_,_,_ in floors])
p(f'Общая длина Cat6: {copper:.1f} м. Закупка с запасом 20%: {copper*1.2:.1f} м; {spools} бухт × 305 м = {spools*305} м. Максимальная постоянная линия — {max(e["length_m"] for e in ends):.1f} м; два патч-корда по 2 м дают канал менее 100 м. Дополнительно 7 патч-кордов: 4 подключения коммутатор–маршрутизатор, 2 межкоммутаторных соединения, 1 связь R-C1–R-EDGE внутри стойки. Патч-корды закупаются отдельно от постоянного кабеля. Концы операторских линий — в зоне ответственности оператора.')
p('Межэтажная оптика: от шкафа C1 до стояка 3 м + подъём 3,6 м + от стояка до шкафа C2 3 м + сервисные петли 2,4 м = 12 м. С 20% запасом 14,4 м → закупка 15 м OM3. Два SFP, два кросса, два патч-корда. Лоток: 36 м × 4 + 20% = 172,8 → 173 м. Ответвления в кабинеты защищаются локальными коробами; их точная смета требует монтажного обследования и в текущую оценку лотка не включена.')
for code,_,_,_ in floors:
 h(f'Постоянные линии {code}',2)
 table(['Устройство','Отдел','(x; y), м','Cat6, м'],[[e['id'],e['department'],f"{e['x']}; {e['y']}",e['length_m']] for e in ends if e['floor']==code])
h('6. Смета оборудования и кабеля (пункт 12)')
p('В смете приняты ориентировочные цены в рублях за единицу оборудования. Расчёт предназначен для предварительной оценки капитальных затрат. Итог включает оборудование и кабельные материалы; монтаж, услуги оператора, лицензии и отдельное резервное хранилище в расчёт не входят.')
table(['Позиция','Характеристики','Кол-во','Цена, ₽','Сумма, ₽'],[[a,b,n,p,n*p] for a,b,n,p in equipment])
p(f'Итого капитальные затраты по данной оценке: {budget_text} ₽. Стоимость ежемесячных услуг связи и программных лицензий учитывается отдельно от капитальных затрат.')
h('7. Логическая адресация и защита')
p('Этажи C1/C2/B1/B2: номера площадок 1/2/11/12. Подсеть VLAN v: 10.s.v.0/24, шлюз .1. VLAN 10 — медицинские рабочие места и МФУ; VLAN 20 — регистратура и администрация в центре; VLAN 30 — серверы только C1; VLAN 90 — гости; VLAN 99 — управление. Гостевые VLAN запрещают доступ к 10.0.0.0/8, разрешают внешний DNS/Интернет. Гостевая точка в учебной модели полностью подключена к VLAN 90, управление через гостевой сегмент не моделируется. В промышленном проекте её управление должно быть отдельно в VLAN 99.')
p('SRV-MIS 10.1.30.10 (HTTP); SRV-MAIL 10.1.30.11 (DNS, EMAIL); SRV-FILE 10.1.30.12 (HTTP как имитация файлового сервиса). Между этажами 10.255.0.0/30; C1–EDGE 10.255.0.4/30; WAN филиалов 10.255.1.0/30 и 10.255.2.0/30. Интернет EDGE–ISP 198.51.100.0/30, внешний сервер 203.0.113.10/24. Документационные диапазоны не обозначают реальных публичных адресов. Статическая маршрутизация; NAT на EDGE.')
p('Реальная медицинская сеть требует разграничения доступа по ролям, шифрования операторских каналов, журналирования, отказоустойчивости и отдельного резервного хранения. В данной предварительной работе не заявляется соответствие требованиям защиты персональных данных; необходимо отдельное обследование и проект защиты. Учебная модель не включает промышленный VPN; её проверки относятся к имитируемым IP-службам, а не к защите реальной медицинской информационной системы.')
h('8. Проверка и ограничения')
p('Согласованы число устройств, портовые назначения, длины трасс, расчёты нагрузки и суммы сметы. Логическая модель clinic8.pkt создана и сохранена в Cisco Packet Tracer 8.2.2.0400. Файл содержит 84 ПК, 8 принтеров, 4 AP, 4 сервера (один внешний), 6 маршрутизаторов, 7 коммутаторов, 4 гостевых планшета и 3 автоматически созданных устройства питания; 112 кабелей. Проверены связь между площадками, доступ к серверам, службы DNS, HTTP и EMAIL, а также изоляция гостевых сегментов. Общий вид модели представлен в пункте 13 презентации. На Wireless0 всех четырёх гостевых планшетов включён DHCP.')
table(['Проверка','Подтверждённый результат'],[['DHCP рабочих ПК','C1: 10.1.20.52/24, шлюз 10.1.20.1; B1: 10.11.10.58/24, шлюз 10.11.10.1; C2: адрес сети 10.2.10.0/24'],['ICMP между площадками','C2/B1/B2 достигают сервера МИС и внешнего сервера; B2 достигает ПК B1. Для B1 повторные ping: 4/4 ответа'],['DNS и HTTP','С B1 открыты mis.clinic8.test, files.clinic8.test и www.example.test'],['EMAIL','Письма B1 → C1 и C2 → B2 доставлены'],['Гостевая изоляция','На C1/C2/B1/B2 внешний HTTP доступен; ICMP и HTTP к серверу 10.1.30.10 заблокированы'],['Повторное открытие','После повторного открытия проекта проверено восстановление доступа к сайтам'],['Границы проверок','В Packet Tracer проверяются учебные IP-службы; нагрузка приложений и скорость оборудования определены расчётным способом']])
h('9. Источники')
p('1. IEEE 802.3 — стандарты Ethernet (параметры 1000BASE-T/1000BASE-SX).\n2. Cisco Business 350 Series Managed Switches Data Sheet: https://www.cisco.com/c/en/us/products/collateral/switches/business-350-series-managed-switches/nb-06-bus350-series-ds-cte-en.html\n3. Cisco 1000 Series Integrated Services Routers Data Sheet: https://www.cisco.com/c/en/us/products/collateral/routers/1000-series-integrated-services-routers-isr/datasheet-c78-739512.html\n4. TP-Link EAP610: https://www.tp-link.com/en/business-networking/omada-wifi-ceiling-mount/eap610/\nТехнические паспорта используются как справочные материалы по выбранным классам оборудования. Цены в смете — ориентировочные.')
D.save(ROOT/'docs/Расчёты_вариант_8.docx')
# slides each mandatory item separate, room plans extra
TEMPLATE=ROOT/'templates/Бауманский_шаблон.pptx'
R=Presentation(TEMPLATE)
COVER_PROTOTYPE=R.slides[0]
BODY_PROTOTYPE=R.slides[11]
# Retain the actual template master, layouts, theme and embedded fonts.
# Example slides are removed; no template instructions become project content.
for sid in list(R.slides._sldIdLst):
 R.part.drop_rel(sid.rId);R.slides._sldIdLst.remove(sid)
BLUE=RGBColor(23,109,234)
TEXT=RGBColor(67,67,67)
FAINT=RGBColor(217,217,217)
WHITE=RGBColor(255,255,255)
REGULAR='ALS Sector Regular'
BOLD='ALS Sector Bold'
def label(s,x,y,w,h,value,size,color,bold=False):
 box=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=box.text_frame
 tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=True;tf.text=value
 for paragraph in tf.paragraphs:
  paragraph.font.name=REGULAR;paragraph.font.bold=bold;paragraph.font.size=PPt(size);paragraph.font.color.rgb=color
 return box
def line(s,x,y,w,h=0):
 sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w if w else .012),Inches(h if h else .012))
 sh.fill.solid();sh.fill.fore_color.rgb=FAINT;sh.line.fill.background()
def background(s,prototype):
 original=prototype._element.cSld.bg
 if original is None:return
 bg=deepcopy(original)
 for elem in bg.iter():
  for attr in [qn('r:embed'),qn('r:link')]:
   rid=elem.get(attr)
   if rid:
    rel=prototype.part.rels[rid]
    newrid=s.part.relate_to(rel.target_ref if rel.is_external else rel.target_part,rel.reltype,is_external=rel.is_external)
    elem.set(attr,newrid)
 current=s._element.cSld.bg
 if current is not None:s._element.cSld.remove(current)
 s._element.cSld.insert(0,bg)
def slide(title,body='',pic=None):
 s=R.slides.add_slide(R.slide_layouts[6]);cover=title.startswith('1.')
 background(s,COVER_PROTOTYPE if cover else BODY_PROTOTYPE)
 # Required section title is first shape, independent of the template artwork.
 label(s,.51,1.65 if cover else .57,12.26,1.75 if cover else .95,title,40 if cover else 26,WHITE if cover else TEXT,cover)
 if cover:
  label(s,.51,.41,12.26,.74,'Московский государственный технический университет им. Н. Э. Баумана',20,WHITE)
  box=label(s,.54,3.66,11.95,2.5,body,18.5,WHITE)
  for paragraph in box.text_frame.paragraphs:paragraph.space_after=PPt(11)
  label(s,.51,6.73,10,.35,'Москва · 2026',18,RGBColor(174,205,255))
 elif pic:
  path=ROOT/'diagrams'/f'{pic}.png';pw,ph=Image.open(path).size
  scale=min(12.26/pw,5.42/ph);iw,ih=pw*scale,ph*scale
  s.shapes.add_picture(str(path),Inches((13.333-iw)/2),Inches(1.64+(5.42-ih)/2),width=Inches(iw),height=Inches(ih))
 else:
  items=body.split('\n');rows=math.ceil(len(items)/2);gap=.28;height=(5.22-gap*(rows-1))/rows
  line(s,6.67,1.62,0,5.24)
  for index,value in enumerate(items):
   x=.76 if index%2==0 else 7.04;y=1.73+(index//2)*(height+gap)
   if index>=2:line(s,x,y-.12,5.43)
   important=value.startswith('Оценка капитальных затрат:')
   if important:
    caption,amount=value.split(': ',1)
    label(s,x,y,5.43,.55,caption,20,TEXT)
    label(s,x,y+.58,5.43,.78,amount,32,BLUE,True)
   else:
    box=label(s,x,y,5.43,height,value,19 if len(value)<115 else 17.5,TEXT)
    # Blue emphasis follows the text-block treatment in the supplied template.
    tf=box.text_frame;prefix,sep,rest=value.partition(':')
    if sep and len(prefix)<55:
     tf.clear();paragraph=tf.paragraphs[0]
     first=paragraph.add_run();first.text=prefix+':';first.font.name=REGULAR;first.font.size=PPt(19 if len(value)<115 else 17.5);first.font.color.rgb=BLUE
     last=paragraph.add_run();last.text=rest;last.font.name=REGULAR;last.font.size=first.font.size;last.font.color.rgb=TEXT
 # Footer is supplementary; original university background/theme is retained.
 if not cover:
  label(s,.76,7.18,10,.18,'ИТС · ДЗ № 1 · Городская поликлиника · Вариант 8',9,RGBColor(161,161,163))
  label(s,12.27,7.12,.45,.27,str(len(R.slides)),12,BLUE)
slide('1. Предварительное проектирование сети поликлиники','МГТУ им. Н. Э. Баумана · факультет ИУ · кафедра ИУ-10\nДисциплина: «Информационно-телекоммуникационные системы»\nДомашнее задание № 1 · Вариант 8\nВыполнил: Устинов Илья Максимович · Группа: ИУ10-59\nПроверила: Федорова В. А. · Москва, 2026')
slide('2. Постановка задачи','Спроектировать ЛВС городской поликлиники: центральный корпус и два филиала в разных районах города.\nОбеспечить МИС, обмен исследованиями, почту, Интернет, печать и беспроводной доступ.\nРезультаты: проект, расчёты DOCX, планы этажей и логическая модель Cisco Packet Tracer.\nМежэтажные и межкорпусные связи — через маршрутизаторы.')
slide('3. Здания и размещение на местности',pic='site')
slide('4. Организационная структура и взаимодействие',pic='org')
slide('5. Прикладные задачи','Регистратура: запись на приём и управление расписанием.\nВрачи и сестринская служба: электронные карты, назначения, результаты.\nДиагностика: исследования DICOM, передача в архив PACS.\nАдминистрация: отчётность; ИТ: обслуживание и резервное копирование.\nВсе площадки: почта, Интернет, печать; посетители: изолированный Wi-Fi.')
slide('6. Программное обеспечение','84 ПК: Windows 11 Pro / совместимая Linux, браузер, офисный пакет, почта, защита рабочих станций.\nДиагностика: DICOM-просмотрщик.\n3 сервера: МИС и PostgreSQL; DNS и почта; файлы, PACS и резервное копирование.\nВ Packet Tracer: только учебная имитация HTTP, DNS и EMAIL; МИС/PACS не развёртываются.')
slide('7. Максимальная нагрузка','Часы наибольшей нагрузки: 08:00–11:00 и 16:00–18:00.\nC = 1,25 × 1,20 × ΣR; 25% служебный запас, 20% рост.\nC1: 180,6; C2: 132,6; каждый филиал: 92,7 Мбит/с.\nСерверные потоки: 375,6 Мбит/с; стык C1–EDGE: 351 Мбит/с.\nИнтернет: 123 Мбит/с → канал 200 Мбит/с.\nПодробные исходные значения и расчёты: «Расчёты_вариант_8.docx».')
slide('8. Стандарт, кабель и топология','Иерархическая звезда с отдельными этажными сегментами.\nРабочие подключения: 1000BASE-T, Cat6 U/UTP LSZH.\nМежэтажная магистраль: 1000BASE-SX, OM3, 15 м к закупке.\nУдалённые филиалы: операторские Ethernet-каналы по 100 Мбит/с.\nНа этажах — коммутаторы; между этажами и зданиями — маршрутизаторы.\nУчебный 2960: 100 Мбит/с к ПК, 1 Гбит/с на магистрали.')
slide('9. Количество и характеристики оборудования','84 ПК: 16 ГБ, SSD 512 ГБ, 1GbE. 3 сервера: 64 ГБ ECC, зеркальные SSD.\n5 маршрутизаторов: 4 этажных и 1 пограничный.\n6 коммутаторов по 24 порта: 144 порта, 99 подключений.\n4 точки Wi-Fi; 8 сетевых МФУ; 4 стойки; 7 ИБП.\nОриентиры: Cisco C1111-8P, CBS350-24P-4G, TP-Link EAP610.\nМодельные ориентиры и ориентировочная стоимость приведены в расчётном приложении.')
for k,(code,title,_,_) in enumerate(floors):slide('10. Структурная схема: '+title,pic=f'floor-{code}')
slide('11. Расчёт кабеля',f'Индивидуальные трассы вычислены по координатам планов, через коридорный лоток.\n99 линий Cat6: {copper:.1f} м; с 20% запасом: {copper*1.2:.1f} м.\nЗакупка: {spools} бухт × 305 м = {spools*305} м.\nМежэтажная оптика: 12 м трасса → 15 м с запасом.\nМаксимальная линия Cat6: {max(e["length_m"] for e in ends):.1f} м; ограничения 90/100 м соблюдены.\nПодробная ведомость: DOCX и cable-routes.csv.')
slide('12. Смета оборудования и кабеля',f'Оценка капитальных затрат: {budget_text} ₽.\nПолная таблица количества, единичных цен и стоимости: DOCX.\nЦены приняты ориентировочно для оценки бюджета проекта.\nНе включены монтаж, лицензии, оплата каналов оператора и отдельное резервное хранилище.\nПолная смета с количеством и ценой каждой позиции приведена в DOCX.')
slide('13. Логическая схема Cisco Packet Tracer' if (ROOT/'diagrams/packet-tracer.png').exists() else '13. Логическая схема — проектный эскиз',pic='packet-tracer' if (ROOT/'diagrams/packet-tracer.png').exists() else 'logical-draft')
slide('13. Файл Cisco и результаты моделирования','Модель clinic8.pkt · Cisco Packet Tracer 8.2.2.0400.\n84 ПК, 4 гостевых планшета, 112 кабелей; сохранены VLAN и WPA2/AES.\nПроверена связь C2/B1/B2 с МИС и внешним сервером; B2 → B1.\nС B1 открыты три сайта по DNS-именам; доставлена почта B1 → C1 и C2 → B2.\nНа всех четырёх гостевых площадках внешний HTTP доступен, внутренний ICMP/HTTP заблокирован.\nНа всех четырёх гостевых планшетах сохранён DHCP. Общий вид модели приведён на предыдущем слайде.')
slide('14. Заключение','Разработан проект сети трёх площадок: 112 сотрудников, 84 ПК, 99 подключений.\nРассчитаны пиковая нагрузка, кабельные трассы и предварительная смета.\nСогласованы этажная маршрутизация, гостевые сегменты и серверные службы.\nВ модели проверены связь между площадками, серверные службы и гостевая изоляция.\nОриентировочная стоимость оборудования и кабельных материалов: 9 061 750 ₽.')
R.save(ROOT/'docs/Презентация_вариант_8.pptx')
print(json.dumps({'endpoints':len(ends),'copper_m':copper,'spools':spools,'budget':total,'slides':len(R.slides)},ensure_ascii=False))
